#!/usr/bin/env python3
"""
Sıfırdan firma kurulumu gibi — baştan sona sıralı manuel test / kurulum listesi (Word).
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "MANUEL_TEST_EKRAN_LISTESI.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    if color:
        run.font.color.rgb = RGBColor(*color)


# Her adım: (başlık, ekran route, yapılacak iş, beklenen sonuç)
# Sıra: bir firmaya sistemi kuruyormuş gibi uçtan uca

PHASES: list[tuple[str, str, list[tuple[str, str, str, str]]]] = [
    (
        "A — Giriş ve firma iskeleti",
        "Önce sisteme girip firmayı / iş yerini / depoyu hazırlayın.",
        [
            (
                "Uygulamayı açıp giriş yapın",
                "# (login)",
                "Admin (veya yetkili) kullanıcı ile oturum açın.",
                "Ana menü ve modüller görünür.",
            ),
            (
                "Aktif firmayı seçin / oluşturun",
                "#companies",
                "TEKS (veya yeni firma) kartını açın; yoksa Yeni ile firma kodu + unvan kaydedin.",
                "Firma listede; üst barda aktif firma doğru.",
            ),
            (
                "İş yeri (plant) tanımlayın",
                "#is-yerleri",
                "Firmaya bağlı en az bir iş yeri ekleyin (kod, ad, şehir).",
                "İş yeri kaydı oluştu; firmaya bağlı.",
            ),
            (
                "Depo tanımlayın",
                "#warehouses",
                "İş yerine bağlı ana depo + gerekirse hammadde/mamul depo ekleyin.",
                "Depolar listelenir; varsayılan depo seçilebilir.",
            ),
            (
                "Kullanıcı–firma yetkisi",
                "#kullanicilar",
                "Test kullanıcısının bu firmaya erişimi olduğundan emin olun.",
                "Firma değiştirince ekranlar dolu / boş doğru davranır.",
            ),
        ],
    ),
    (
        "B — Sabit tanımlar (bir kez)",
        "Belge ve kartlar için ortak parametreler. Eksikse sonraki adımlarda lookup boş kalır.",
        [
            (
                "Numaratörler",
                "#numarator-tanimi",
                "Stok, cari, sipariş, irsaliye, fatura, iş emri serilerinin aktif olduğunu kontrol edin.",
                "Yeni kayıtta kod otomatik gelir.",
            ),
            (
                "Ölçü birimleri",
                "#olcu-birimleri",
                "AD, KG, MT, LT vb. temel birimleri ekleyin / doğrulayın.",
                "Stok kartında birim seçilebilir.",
            ),
            (
                "Para birimi ve döviz kuru",
                "#para-birimi-tanimi → #doviz-exchange_rate-tanimi",
                "TRY (ve gerekirse USD/EUR) + güncel kur girin.",
                "Belgelerde PB ve kur doluyor.",
            ),
            (
                "Ülke / şehir / ilçe",
                "#country-tanimi → #sehir-tanimi → #district-tanimi",
                "TR ve ilgili şehirleri tamamlayın.",
                "Cari adres lookup’ları dolu.",
            ),
            (
                "Ödeme ve teslimat koşulları",
                "#odeme-kosulu-tanimi → #teslimat-kosulu-tanimi",
                "Peşin, 30 gün vb. ve teslimat şartlarını ekleyin.",
                "Sipariş / faturada seçilebiliyor.",
            ),
            (
                "Hareket kodları ve belge tipleri",
                "#hareket-kodlari → #belge-tipleri",
                "Stok giriş/çıkış kodları ve belge tiplerini kontrol edin (stok etkisi).",
                "Mal kabul / irsaliye onayında stok hareketi oluşur.",
            ),
            (
                "Kategori / dil / banka (isteğe bağlı)",
                "#kategori-tanimi, #dil-tanimi, #banka-tanimi",
                "İhtiyaç kadar tanımlayın.",
                "Lookup’larda görünür.",
            ),
        ],
    ),
    (
        "C — Ana kartlar (stok + cari)",
        "İş süreçlerinin dayandığı kartlar. Önce hammadde/mamul, sonra tedarikçi/müşteri.",
        [
            (
                "Hammadde stok kartları",
                "#stok-karti",
                "En az 2–3 hammadde: kod, ad, birim, takip birimi, tip=Hammadde.",
                "Liste ve detay kayıtlı; lookup’ta çıkar.",
            ),
            (
                "Mamul stok kartı",
                "#stok-karti",
                "1 mamul kartı açın (tip=Mamul), birim ve fiyat alanlarını doldurun.",
                "Mamul kartı kaydedildi.",
            ),
            (
                "Stok–depo / iş yeri bağlantısı",
                "#stok-depo (veya stok kartı altı)",
                "Kartların depo/iş yeri ilişkilerini doğrulayın.",
                "Bakiye ekranında stok görünür (miktar 0 olabilir).",
            ),
            (
                "Tedarikçi cari",
                "#cari-kartlari",
                "En az 1 tedarikçi: kod, unvan, tip, ödeme koşulu, PB.",
                "Satınalma belgelerinde cari seçilir.",
            ),
            (
                "Müşteri cari",
                "#cari-kartlari",
                "En az 1 müşteri cari açın.",
                "Satış belgelerinde cari seçilir.",
            ),
        ],
    ),
    (
        "D — Üretim tanımı (BOM + rota)",
        "Mamulü üretmek için ağaç ve operasyon sırası.",
        [
            (
                "İş merkezi ve iş istasyonu",
                "#is-merkezleri → #is-istasyonlari",
                "En az bir iş merkezi + istasyon tanımlayın.",
                "Rotada seçilebilir.",
            ),
            (
                "Ürün ağacı (BOM)",
                "#urun-agaci",
                "Mamul altına hammaddeleri miktarlarıyla bağlayın; kaydedin / patlatmayı deneyin.",
                "BOM satırları doğru; patlatma / seviye listesi çalışır.",
            ),
            (
                "Rota (operasyonlar)",
                "#routings",
                "OP10, OP20… sırasıyla operasyon ekleyin; iş merkezi bağlayın.",
                "İş emrine rota uygulanabiliyor.",
            ),
            (
                "Üretim emri tipi (varsa)",
                "#uretim-emri-tipleri",
                "Standart tipi kontrol edin.",
                "İş emri formunda tip seçilir.",
            ),
        ],
    ),
    (
        "E — Satınalma döngüsü (alış)",
        "Hammaddeyi tedarikçiden alıp depoya sokun.",
        [
            (
                "Satınalma talebi",
                "#satin-alma-talepleri",
                "Hammadde satırlarıyla talep oluşturun; kaydedin.",
                "Talep numarası üretildi, durum uygun.",
            ),
            (
                "Teklif isteği / satınalma teklifi (isteğe bağlı)",
                "#satin-alma-teklif-istekleri → #satin-alma-teklifleri",
                "Talepten teklife ilerleyin veya doğrudan teklif girin.",
                "Teklif satırları ve tedarikçi bağlı.",
            ),
            (
                "Satınalma siparişi",
                "#satin-alma-siparisleri",
                "Tekliften veya sıfırdan SA siparişi oluşturun; onaylayın (onay varsa).",
                "Sipariş Onaylı / açık durumda.",
            ),
            (
                "Mal kabul",
                "#mal-kabuller",
                "SA siparişinden mal kabul oluşturun; miktarları kabul edin; onaylayın.",
                "Stok bakiyesi arttı; hareket oluştu.",
            ),
            (
                "Tedarikçi irsaliyesi (varsa)",
                "#tedarikci-irsaliyeleri",
                "Gelen irsaliyeyi kaydedin / ilişkilendirin.",
                "Belge listede görünür.",
            ),
            (
                "Alış faturası",
                "#fatura-yonetimi",
                "Tür=Alış; mal kabulden fatura oluşturun; kaydedin / muhasebeleştirin.",
                "Fatura oluştu; 3-way ekranında eşleşme görülebilir.",
            ),
            (
                "Fatura kontrolü (3-way)",
                "#fatura-kontrolleri",
                "Sipariş–mal kabul–fatura eşlemesini kontrol edin.",
                "Sonuç Bekliyor değilse tutarlı / fark açıklanabilir.",
            ),
        ],
    ),
    (
        "F — Üretim döngüsü",
        "Depodaki hammadde ile mamul üretin.",
        [
            (
                "Üretim planı (isteğe bağlı)",
                "#uretim-planlari",
                "Plan kaydı açın veya doğrudan iş emrine geçin.",
                "Plan görünür veya atlandı.",
            ),
            (
                "İş emri oluştur",
                "#uretim-emirleri",
                "Mamul + miktar; BOM/rota uygula; kaydet.",
                "Bileşen ve operasyon satırları doldu.",
            ),
            (
                "İş emrini yürüt / üretim girişi",
                "#uretim-emirleri → #uretim-girisleri",
                "Operasyonları tamamlayın; mamul girişini onaylayın.",
                "Mamul stok bakiyesi arttı; hammadde düştü (tüketim varsa).",
            ),
            (
                "Depo bakiyesi kontrolü",
                "#depo-bakiye",
                "Hammadde ve mamul bakiyelerini doğrulayın.",
                "Miktarlar beklenenle uyumlu.",
            ),
        ],
    ),
    (
        "G — Satış döngüsü",
        "Mamulü müşteriye sevk edip faturalayın.",
        [
            (
                "Satış teklifi",
                "#satis-teklifleri",
                "Müşteri + mamul satırı ile teklif oluşturun.",
                "Teklif kaydı ve tutar doğru.",
            ),
            (
                "Satış siparişi",
                "#satis-siparisleri",
                "Tekliften siparişe çevirin veya yeni sipariş açın; onaylayın.",
                "Sipariş açık/onaylı.",
            ),
            (
                "Sevk emri",
                "#sevk-emirleri",
                "Siparişten sevk oluşturun; depo seçin.",
                "Sevk satırları dolu.",
            ),
            (
                "Satış irsaliyesi",
                "#irsaliye-yonetimi",
                "Sevkten irsaliye; onayla (stok çıkış).",
                "Mamul bakiyesi azaldı; e-belge durumu (simülasyon) güncellendi.",
            ),
            (
                "Satış faturası",
                "#fatura-yonetimi",
                "Tür=Satış; irsaliyeden fatura oluşturun; kaydedin.",
                "Fatura ve cari borç / muhasebe etkisi oluştu.",
            ),
            (
                "Cari ekstre",
                "#cari-ekstre",
                "Müşteri ve tedarikçi hareketlerini kontrol edin.",
                "Alış/satış faturaları ekstrede görünür.",
            ),
        ],
    ),
    (
        "H — Kalite (örnek)",
        "Mal kabul veya üretim sonrası kalite emri.",
        [
            (
                "Kontrol planı / ret nedeni",
                "#kk-kontrol-planlari → #kk-ret-nedenleri",
                "Temel tanımları doğrulayın.",
                "Emir formunda seçilebilir.",
            ),
            (
                "Giriş kalite emri (GKK)",
                "#kk-emirleri-gkk",
                "Mal kabule bağlı veya bağımsız GKK emri açın; sonuç girin.",
                "Emir kapandı / uygunluk kaydı var.",
            ),
            (
                "Çıkış kalite (CKK) — isteğe bağlı",
                "#kk-emirleri-ckk",
                "Sevk öncesi örnek kontrol.",
                "Kayıt oluştu.",
            ),
        ],
    ),
    (
        "I — CRM (isteğe bağlı)",
        "Satış öncesi pipeline.",
        [
            (
                "Aday → kontak → fırsat",
                "#crm-adaylar → #crm-kontaklar → #crm-firsatlar",
                "Bir fırsat oluşturup aşama ilerletin; teklife bağlayın.",
                "Fırsat teklife dönüştürülebilir.",
            ),
            (
                "Faaliyet / görev / ajanda",
                "#crm-faaliyetler → #crm-gorevler → #crm-ajanda",
                "Bir görüşme veya görev kaydı girin.",
                "Ajandada görünür.",
            ),
        ],
    ),
    (
        "J — Onay ve finans iskeleti",
        "Onay rotaları ve muhasebe temel tanımları.",
        [
            (
                "Onay rol / kural / rota",
                "#onay-role-tanim → #onay-kural-tanim → #onay-rota-tanim",
                "SA siparişi veya fatura için basit rota tanımlayın.",
                "Belgede Onayla akışı çalışır; #onaylarim’da görünür.",
            ),
            (
                "Hesap planı + mali yıl + dönem",
                "#hesap-plani → #mali-yillar → #muhasebe-donemleri",
                "Temel hesaplar, açık mali yıl ve dönemler.",
                "Fiş / fatura muhasebeleştirmesi dönem açıkken çalışır.",
            ),
            (
                "Vergi kodu / ülke paketi (TR)",
                "#finans-vergi-kodlari → #finans-ulke-paketleri",
                "KDV kodlarını ve TR paketini doğrulayın.",
                "Faturada vergi hesaplanır.",
            ),
            (
                "Banka / kasa",
                "#finans-banka-hesaplari → #finans-kasa-hesaplari",
                "Bir banka ve bir kasa hesabı açın.",
                "Nakit hareketinde seçilir.",
            ),
            (
                "Tahsilat / ödeme örneği",
                "#tahsilat-odemeler veya #finans-nakit-hareketleri",
                "Müşteriden tahsilat veya tedarikçiye ödeme girin.",
                "Cari bakiye ve kasa/banka güncellendi.",
            ),
            (
                "Mizan / yevmiye özeti",
                "#mizan → #finans-yevmiye",
                "Dönem aralığı ile rapor alın.",
                "Hareketli hesaplar listelenir.",
            ),
        ],
    ),
    (
        "K — İnsan kaynakları (temel)",
        "Firma için İK iskeleti.",
        [
            (
                "Departman ve pozisyon",
                "#hr-departmanlar → #hr-pozisyonlar",
                "Organizasyon ağacını kurun.",
                "Personel kartında seçilir.",
            ),
            (
                "Çalışma takvimi / vardiya / tatil",
                "#hr-calisma-takvimleri → #hr-vardiyalar → #hr-resmi-tatiller",
                "Standart takvim ve bir vardiya tanımlayın.",
                "Personelde atanabilir.",
            ),
            (
                "Personel kartı",
                "#hr-personel",
                "En az 1 personel: sicil, ad soyad, işe giriş, departman.",
                "Kart kaydedildi.",
            ),
            (
                "İzin türü + izin talebi (örnek)",
                "#hr-izin-turleri → #hr-izin-talepleri",
                "Yıllık izin türü ve bir talep oluşturun; gönder/onayla.",
                "Bakiye / talep durumu güncellendi.",
            ),
        ],
    ),
    (
        "L — Sistem kontrolü ve kapanış",
        "Kurulumun tutarlı olduğunu doğrulayın.",
        [
            (
                "Kokpit",
                "#kokpit",
                "Açık sipariş, sevk, kalite özetine bakın.",
                "TEKS verisi özetlerde görünür.",
            ),
            (
                "Audit log",
                "#audit-log",
                "Son oluşturduğunuz belgelerden birinin logunu filtreleyin.",
                "Kullanıcı + tablo + kayıt izi var.",
            ),
            (
                "Tablo Bilgileri",
                "#tablo-bilgileri",
                "İnsan Kaynakları ve Finans gruplarının listelendiğini kontrol edin.",
                "Ekran → ana tablo eşlemesi açılıyor.",
            ),
            (
                "Firma izolasyonu",
                "Firma seçici",
                "Başka firmaya geçip TEKS kayıtlarının görünmediğini doğrulayın; geri dönün.",
                "Veri firmalar arası karışmaz.",
            ),
        ],
    ),
]


def add_para(doc, text, *, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)

    h = doc.add_heading("KTB ERP — Sıfırdan Firma Kurulum / Manuel Test Listesi", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT

    now = datetime.now(ZoneInfo("Europe/Istanbul")).strftime("%d.%m.%Y %H:%M")
    add_para(
        doc,
        f"Oluşturma (Europe/Istanbul): {now}\n"
        "Nasıl kullanılır: Madde madde yukarıdan aşağı gidin. Her adımı bitirmeden sonrakine geçmeyin.\n"
        "Hedef firma örneği: TEKS — Denim Tekstil Üretim A.Ş. (yeni firma ile de aynı sıra).\n"
        "Her maddede: (1) ekranı açın (2) işlemi yapın (3) beklenen sonucu işaretleyin.",
        size=10,
        space_after=10,
    )

    add_para(doc, "İşaretleme:  ☐ Yapılmadı    ☑ Tamam    ✗ Hata (not yazın)", bold=True, size=10)

    step_no = 0
    for phase_title, phase_intro, steps in PHASES:
        doc.add_heading(phase_title, level=1)
        add_para(doc, phase_intro, size=10, space_after=8)

        for title, screen, action, expected in steps:
            step_no += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(f"{step_no}. {title}")
            set_run_font(r, size=12, bold=True, color=(0x1F, 0x3A, 0x5F))

            add_para(doc, f"Ekran: {screen}", size=10, space_after=2)
            add_para(doc, f"Yapılacak: {action}", size=10, space_after=2)
            add_para(doc, f"Beklenen: {expected}", size=10, space_after=2)
            add_para(doc, "Sonuç:  ☐ Tamam    ☐ Hata — not: ___________________________", size=10, space_after=6)

    doc.add_heading("Bitince", level=1)
    for t in [
        "Tüm zorunlu adımlar (A–G ve L) tamam mı?",
        "Kritik hataları ekran + adım no ile not ettiniz mi?",
        "TEKS iş verisini tekrar temizleyecekseniz: node scripts/purge-firma-veri.js TEKS",
    ]:
        doc.add_paragraph(t, style="List Number")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT} ({step_no} adım, {len(PHASES)} faz)")


if __name__ == "__main__":
    main()
