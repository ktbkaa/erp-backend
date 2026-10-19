#!/usr/bin/env python3
"""MRP Kullanım Kılavuzu (Word) üretir."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "26_MRP_KULLANIM_KILAVUZU.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size={1: 18, 2: 14, 3: 12}.get(level, 11), bold=True)
    return p


def add_para(doc, text, *, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        set_run_font(run, size=11)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, size=10, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    # Kapak
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MRP Kullanım Kılavuzu")
    set_run_font(r, size=26, bold=True, color=RGBColor(0x1A, 0x3A, 0x5C))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Malzeme İhtiyaç Planlaması — Süreç, Ekranlar ve Örnek Senaryo")
    set_run_font(r, size=12, color=RGBColor(0x44, 0x55, 0x66))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"ERP Demo · Migration 028_mrp_surec · {date.today().isoformat()}\n"
        "Kaynaklar: klasik MRP, Canias «MRP - TR - Kullanıcı Seviyesi» (33 video), "
        "Odoo 17 MRP oynatma listesi (15 video)"
    )
    set_run_font(r, size=9, color=RGBColor(0x66, 0x77, 0x88))

    doc.add_paragraph()

    # 1
    add_heading(doc, "1. MRP nedir?", 1)
    add_para(
        doc,
        "MRP (Material Requirements Planning / Malzeme İhtiyaç Planlaması), bağımsız talebi "
        "(satış siparişi, tahmin, bütçe, manuel ihtiyaç) ürün ağacı (BOM) üzerinden bileşenlere "
        "patlatır; eldeki stok, açık satınalma ve açık üretim emirlerini düşerek net ihtiyacı "
        "hesaplar; satınalma veya üretim önerisi üretir.",
    )
    add_para(doc, "Klasik formül (özet):", bold=True)
    add_bullets(
        doc,
        [
            "Brüt ihtiyaç = Bağımsız ihtiyaç + Bağımlı ihtiyaç (BOM patlatma)",
            "Net ihtiyaç = max(0, Brüt + Emniyet stok + ROP boşluğu − Eldeki − Açık SA − Açık UE)",
            "Öneri miktar = Lot kuralı (lot-for-lot veya sabit sipariş miktarı)",
            "Öneri tipi = Tedarik tipi «Üretim» → Üretim emri; aksi → Satınalma",
            "Plan başlangıç = İhtiyaç tarihi − Temin süresi (gün)",
        ],
    )

    # 2
    add_heading(doc, "2. Araştırma özeti (Canias + Odoo)", 1)
    add_heading(doc, "2.1 Canias — MRP - TR - Kullanıcı Seviyesi (33 video)", 2)
    add_para(
        doc,
        "Oynatma listesi (Canias ERP YouTube): kullanıcı seviyesinde MRP parametreleri, "
        "bağımsız ihtiyaç girişi, çalıştırma (MRPT*), sonuç listesi, önerilerin satınalma/"
        "üretim belgelerine dönüştürülmesi, lead time / emniyet stok / lot büyüklüğü kavramları.",
    )
    add_bullets(
        doc,
        [
            "Bağımsız ihtiyaç (forecast / sipariş / manuel) planlamanın girdisidir.",
            "Ürün ağacı ve tedarik tipi (satınalma / üretim) öneri türünü belirler.",
            "Çalışma kaydı ile sonuçlar versiyonlanır; aynı dönemde tekrar çalıştırılabilir.",
            "Net ihtiyaç satırlarından satınalma talebi / üretim emri oluşturulur.",
            "Temin süresi ile geri çizelgeleme (ihtiyacı karşılamak için ne zaman başlatılmalı) yapılır.",
        ],
    )

    add_heading(doc, "2.2 Odoo 17: MRP (15 video — Ahmed Elsaka)", 2)
    add_para(
        doc,
        "Oynatma listesi: BOM, Manufacturing Order (MO), reordering rules, MRP run, "
        "hammadde rezervasyonu ve üretimden mamule akış. Bizim modele yansıyanlar:",
    )
    add_bullets(
        doc,
        [
            "BOM (ürün ağacı) → bağımlı ihtiyaç patlatması",
            "Manufacturing Order → üretim emri dönüşümü",
            "Reordering / safety stock → emniyet stok + yeniden sipariş seviyesi (ROP)",
            "Purchase önerisi → satınalma talebi",
            "Lead time → stok kartındaki temin_suresi_gun",
        ],
    )

    add_heading(doc, "2.3 Bizim sadeleştirilmiş model", 2)
    add_para(
        doc,
        "Canias’ın tam MRPT* ekran seti ve Odoo’nun kapasite/MO planlayıcısı yerine, "
        "demo ERP’de tek motor + dört ekran + belge dönüşüm butonları kullanılır. "
        "Kapasite planlama (CRP), alternatif BOM ve ileri seviye lot politikaları sonraki dalgaya bırakılmıştır.",
    )

    # 3
    add_heading(doc, "3. Süreç modeli", 1)
    add_numbered(
        doc,
        [
            "Stok kartında planlama parametrelerini ayarla (temin süresi, emniyet, ROP, sabit lot).",
            "Ürün ağacı (BOM) ve rota tanımlı olsun (üretim tipi mamuller için).",
            "Bağımsız ihtiyaç gir (veya açık satış siparişi zaten talep oluştursun).",
            "MRP Çalıştır → sonuçlar kayda alınır (mrp_calismalari + mrp_sonuclar).",
            "Net>0 satırlardan SA Talep ve/veya Üretim Emri oluştur.",
            "SA / UE süreçlerini normal belgelerle tamamla; bir sonraki MRP’de açık arz düşülür.",
        ],
    )

    add_heading(doc, "3.1 Veri tabloları", 2)
    add_table(
        doc,
        ["Tablo", "Rol"],
        [
            ["mrp_bagimsiz_ihtiyaclar", "Manuel / tahmin / bütçe bağımsız ihtiyaç"],
            ["mrp_calismalari", "Çalışma üst bilgisi (no, durum, satır/net sayısı)"],
            ["mrp_sonuclar", "Çalışmaya bağlı brüt/net/öneri satırları"],
            ["stoklar.temin_suresi_gun", "Lead time (gün)"],
            ["stoklar.yeniden_siparis_seviyesi", "ROP"],
            ["stoklar.sabit_siparis_miktari", "Sabit lot (null = lot-for-lot)"],
            ["stoklar.emniyet_stok", "Emniyet stoku (nete eklenir)"],
        ],
    )

    add_heading(doc, "3.2 Talep ve arz kaynakları", 2)
    add_table(
        doc,
        ["Kaynak", "Tür", "Nasıl kullanılır"],
        [
            ["mrp_bagimsiz_ihtiyaclar (aktif)", "Talep", "Bağımsız ihtiyaç"],
            ["Açık satış siparişi kalanı", "Talep", "Bağımsız ihtiyaç"],
            ["Açık UE bileşenleri", "Talep", "Bağımlı ihtiyaç (çift patlatma yok)"],
            ["BOM patlatma", "Talep", "Üretim tipi bağımsız → bileşen bağımlı"],
            ["Depo bakiyesi (Kullanılabilir)", "Arz", "Eldeki stok"],
            ["Açık satınalma siparişi", "Arz", "Açık SA"],
            ["Açık üretim emri kalanı", "Arz", "Planlanan üretim"],
        ],
    )

    # 4
    add_heading(doc, "4. Ekranlar ve bağlantılar", 1)
    add_para(doc, "Menü yolu: Stok → Planlama", bold=True)
    add_table(
        doc,
        ["Ekran", "Kod", "Ne yapar", "Bağlantılar"],
        [
            [
                "MRP Bağımsız İhtiyaç",
                "mrp-bagimsiz-ihtiyac",
                "BI CRUD (numaratör MRP_BI)",
                "Stok lookup; firma / iş yeri",
            ],
            [
                "MRP Çalıştır",
                "mrp-calisma",
                "Çalışma geçmişi; «MRP Çalıştır» butonu",
                "İncele → MRP Sonuçları",
            ],
            [
                "MRP Anlık Hesap",
                "mrp-hesap",
                "Kaydetmeden anlık /mrp listesi",
                "SA Talep · Üretim Emri · MRP Çalıştır",
            ],
            [
                "MRP Sonuçları",
                "mrp-sonuclar",
                "Son (veya seçilen) çalışmanın satırları",
                "SA Talep · Üretim Emri → ilgili belgeler",
            ],
            [
                "Stok Kartı",
                "stoklar",
                "Temin / ROP / sabit lot / emniyet",
                "Planlama sekmesi alanları",
            ],
        ],
    )

    add_heading(doc, "4.1 API uçları", 2)
    add_table(
        doc,
        ["Metot", "Yol", "Açıklama"],
        [
            ["GET", "/mrp", "Anlık hesap listesi"],
            ["GET/POST/PUT/DELETE", "/mrp/bagimsiz-ihtiyaclar", "BI CRUD"],
            ["GET", "/mrp/calismalar", "Çalışma listesi"],
            ["GET", "/mrp/calismalar/:id/sonuclar", "Çalışma sonuçları"],
            ["POST", "/mrp/calistir", "Çalıştır ve kaydet"],
            ["POST", "/mrp/talep-olustur", "Satınalma önerilerinden SA talep"],
            ["POST", "/mrp/uretim-emri-olustur", "Üretim önerilerinden UE"],
        ],
    )

    add_heading(doc, "4.2 Eylem butonları", 2)
    add_bullets(
        doc,
        [
            "MRP Çalıştır — hesaplar, mrp_calismalari + mrp_sonuclar yazar, Sonuçlar ekranına gider.",
            "SA Talep — oneri=Satınalma ve net>0 satırlardan satin_alma_talepleri oluşturur; sonuç satırını dönüştürülmüş işaretler.",
            "Üretim Emri — oneri=Üretim emri satırlarından uretim_emirleri + bileşenler (BOM kopyası) oluşturur.",
            "Çalışma listesinde İncele / satır tıklama — ilgili çalışmanın sonuçlarını açar.",
        ],
    )

    # 5
    add_heading(doc, "5. Örnek senaryo (demo)", 1)
    add_para(
        doc,
        "Script: node scripts/demo-mrp-ornek.js — mamul U00001 (Kompanzasyon Rölesi KR-200), "
        "BOM: H00001×1, H00002×4, H00003×2, H00004×1.",
    )
    add_numbered(
        doc,
        [
            "Aktif bağımsız ihtiyaç: 300 adet U00001 (ihtiyac_no BI…, kaynak Manuel, +14 gün).",
            "MRP çalıştır → çalışma MR… kaydı; net>0 satırlar (ör. U00001 üretim, H00001–H00004 satınalma).",
            "SA talep L… — satınalma önerili hammaddeler.",
            "Üretim emri E… — U00001 için planlanan miktar; bileşenler BOM’dan kopyalanır.",
            "UI doğrulama: Planlama → MRP Sonuçları / Bağımsız İhtiyaç; Satınalma Talepleri; Üretim Emirleri.",
        ],
    )
    add_para(
        doc,
        "Not: Eldeki stok ve açık SA yüksekse hammadde net ihtiyacı sıfır olabilir. "
        "Örnek script bilerek 300 adet kullanır ki patlatma sonrası satınalma önerileri görünsün. "
        "Emniyet stoku olan ve bakiyesi olmayan diğer stoklar (ör. H00019) da ROP/emniyet nedeniyle öneriye düşebilir.",
    )

    add_heading(doc, "5.1 Beklenen hesap (özet)", 2)
    add_table(
        doc,
        ["Stok", "Bağımsız", "Bağımlı (≈)", "Öneri"],
        [
            ["U00001", "300 (+ açık SO varsa)", "—", "Üretim emri (eldeki+açık UE düşülür)"],
            ["H00001", "—", "≈300×1", "Satınalma (stok/SA yetersizse)"],
            ["H00002", "—", "≈300×4", "Satınalma"],
            ["H00003", "—", "≈300×2", "Satınalma"],
            ["H00004", "—", "≈300×1", "Satınalma"],
        ],
    )

    # 6
    add_heading(doc, "6. Adım adım kullanım (kullanıcı)", 1)
    add_heading(doc, "6.1 Parametreler", 2)
    add_numbered(
        doc,
        [
            "Stok kartını açın → Planlama alanları: Temin süresi (gün), Emniyet stok, Yeniden sipariş seviyesi, Sabit sipariş miktarı.",
            "Mamul için Tedarik tipi = Üretim; hammadde için Satınalma.",
            "Ürün Ağacı ve Rota tanımlı olsun (üretim emri dönüşümü için zorunlu).",
        ],
    )
    add_heading(doc, "6.2 Bağımsız ihtiyaç", 2)
    add_numbered(
        doc,
        [
            "Stok → Planlama → MRP Bağımsız İhtiyaç → Yeni.",
            "Stok, miktar, ihtiyaç tarihi, kaynak (Manuel / Siparis / Tahmin / Bütçe) doldurun; kaydedin.",
            "Açık satış siparişleri ayrıca otomatik bağımsız ihtiyaç sayılır; çift kayıt gerekmez.",
        ],
    )
    add_heading(doc, "6.3 Çalıştırma ve dönüşüm", 2)
    add_numbered(
        doc,
        [
            "MRP Çalıştır veya MRP Anlık Hesap ekranında «MRP Çalıştır»a basın.",
            "MRP Sonuçları’nda Net ve Öneri kolonlarını inceleyin.",
            "«SA Talep» ile satınalma önerilerini talep belgesine çevirin.",
            "«Üretim Emri» ile üretim önerilerini UE’ye çevirin.",
            "Dönüşen satırlar «Dönüştü» ve belge no ile işaretlenir; aynı çalışmadan mükerrer dönüşüm engellenir.",
        ],
    )

    # 7
    add_heading(doc, "7. Bilinçli sınırlar", 1)
    add_bullets(
        doc,
        [
            "Zaman kovalı (bucket) çok periyotlu MRP yok — tek ufuk / ihtiyaç tarihi odaklı.",
            "Kapasite planlama (CRP), iş merkezi yükü, OEE yok.",
            "Alternatif BOM / varyant / ikame otomatik seçimi yok.",
            "Pegging (hangi üst ihtiyaçtan geldi) sınırlı: kaynak_stok_id gösterim amaçlı.",
            "Firmalar arası izolasyon: firma_id / is_yeri_id filtreleri desteklenir; çok firma demoda aktif firma header’ı kullanılır.",
        ],
    )

    # 8
    add_heading(doc, "8. Teknik referans", 1)
    add_bullets(
        doc,
        [
            "Migration: db/migrations/028_mrp_surec.sql",
            "Motor: src/services/mrpService.js (mrpHesapla, mrpCalistir)",
            "Rotalar: routes/mrp.js",
            "UI: public/app.js (planlama grubu ekranları) + public/index.html (butonlar)",
            "Demo: scripts/demo-mrp-ornek.js",
            "Numaratör: MRP_BI (BI…), MRP_CALISMA (MR…)",
        ],
    )

    add_heading(doc, "9. Hızlı kontrol listesi", 1)
    add_table(
        doc,
        ["Kontrol", "Beklenen"],
        [
            ["npm run migrate", "028_mrp_surec uygulanmış"],
            ["GET /mrp", "JSON liste (net/öneri kolonları)"],
            ["BI kaydı + POST /mrp/calistir", "calisma_no + sonuclar"],
            ["node scripts/demo-mrp-ornek.js", "BI + SA talep + UE logları"],
            ["UI Planlama menüsü", "4 ekran + butonlar görünür"],
        ],
    )

    footer = doc.add_paragraph()
    r = footer.add_run(
        "Bu kılavuz demo ERP MRP v1 içindir. Tasarım iyileştirmeleri sonraki iterasyonda ele alınacaktır."
    )
    set_run_font(r, size=9, color=RGBColor(0x66, 0x77, 0x88))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Yazıldı: {OUT}")


if __name__ == "__main__":
    main()
